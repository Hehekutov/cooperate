from copy import deepcopy
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


TEMPLATE_PATH = Path("/Users/sskaaat/Downloads/testing-template.docx")
OUTPUT_PATH = Path("/Users/sskaaat/DevProjects/cooperate/output/doc/testing-scenarios.docx")


SCENARIOS = [
    {
        "area": "Frontend",
        "priority": "Высокий",
        "title": "Регистрация компании с созданием аккаунта директора",
        "summary": "Проверить, что форма регистрации компании на frontend валидирует обязательные поля и отправляет корректные данные на backend.",
        "steps": [
            "1. Открыть страницу регистрации компании.",
            "2. Заполнить название компании, описание, ФИО директора, должность, телефон и пароль.",
            "3. Нажать кнопку регистрации.",
        ],
        "data": "companyName=Acme, companyDescription=Internal idea platform, directorName=Иван Иванов, directorPosition=Директор, phone=+79990001122, password=secret123",
        "expected": "Форма успешно отправляется, пользователь получает подтверждение регистрации и автоматически переходит в систему под учетной записью директора.",
        "actual": "Проверка не выполнялась.",
        "precondition": "Пользователь не авторизован, открыта форма регистрации.",
        "postcondition": "Создана новая компания, директорский аккаунт доступен в системе.",
        "comments": "Frontend. Проверить отображение ошибок валидации при пропуске обязательных полей.",
    },
    {
        "area": "Frontend",
        "priority": "Высокий",
        "title": "Вход в систему с корректными учетными данными",
        "summary": "Проверить, что экран входа принимает телефон и пароль, а после успешной авторизации открывает рабочую область пользователя.",
        "steps": [
            "1. Открыть экран входа.",
            "2. Ввести телефон и пароль существующего пользователя.",
            "3. Нажать кнопку входа.",
        ],
        "data": "phone=+70000000003, password=employee123",
        "expected": "Пользователь авторизуется, получает доступ к приложению и видит интерфейс согласно своей роли.",
        "actual": "Проверка не выполнялась.",
        "precondition": "В системе есть активная учетная запись сотрудника.",
        "postcondition": "Сессия пользователя активна, токен сохранен на клиенте.",
        "comments": "Frontend. После входа не должно быть повторного показа экрана авторизации до выхода из сессии.",
    },
    {
        "area": "Frontend",
        "priority": "Средний",
        "title": "Отображение меню и данных профиля по роли пользователя",
        "summary": "Проверить, что frontend показывает только доступные действия для директора, администратора и сотрудника.",
        "steps": [
            "1. Авторизоваться под директором.",
            "2. Проверить видимость разделов сотрудников, модерации и решений.",
            "3. Повторить проверку под администратором и сотрудником.",
        ],
        "data": "director=+70000000001/director123, admin=+70000000002/admin123, employee=+70000000003/employee123",
        "expected": "В интерфейсе отображаются только разрешенные разделы и кнопки для текущей роли; недоступные действия скрыты или заблокированы.",
        "actual": "Проверка не выполнялась.",
        "precondition": "В системе заведены пользователи с тремя ролями.",
        "postcondition": "Пользователь остается на главной странице со списком доступных ему разделов.",
        "comments": "Frontend. Проверить, что директор видит управление сотрудниками, а сотрудник не видит административные действия.",
    },
    {
        "area": "Frontend",
        "priority": "Высокий",
        "title": "Создание новой идеи через форму",
        "summary": "Проверить, что форма создания идеи на frontend валидирует заголовок и описание и успешно отправляет запрос на backend.",
        "steps": [
            "1. Открыть раздел создания идеи.",
            "2. Ввести заголовок и описание идеи.",
            "3. Нажать кнопку отправки.",
        ],
        "data": "title=Установить кофемашину, description=Сотрудники тратят время на выход за кофе.",
        "expected": "После отправки идея появляется в списке пользователя и в общей ленте согласно текущему статусу.",
        "actual": "Проверка не выполнялась.",
        "precondition": "Пользователь авторизован как сотрудник или администратор.",
        "postcondition": "Созданная идея доступна в разделе своих идей и ожидает дальнейшего жизненного цикла.",
        "comments": "Frontend. Проверить сообщения об ошибке, если заголовок или описание пустые.",
    },
    {
        "area": "Frontend",
        "priority": "Средний",
        "title": "Переключение между активными идеями и архивом",
        "summary": "Проверить, что frontend корректно переключает scope списка идей и отображает разные наборы записей.",
        "steps": [
            "1. Открыть список идей.",
            "2. Переключиться на активные идеи.",
            "3. Переключиться на архив идей.",
        ],
        "data": "scope=active, scope=archive",
        "expected": "Список обновляется без ошибок, активные идеи и архив показывают разные наборы данных согласно выбранному фильтру.",
        "actual": "Проверка не выполнялась.",
        "precondition": "В системе есть идеи в активном состоянии и в архиве.",
        "postcondition": "На экране отображается выбранный список идей.",
        "comments": "Frontend. Проверить индикатор загрузки и отсутствие дублирования карточек при смене фильтра.",
    },
    {
        "area": "Frontend",
        "priority": "Средний",
        "title": "Просмотр карточки идеи и доступных действий",
        "summary": "Проверить, что на странице идеи отображаются детали, статус и доступные действия в зависимости от роли пользователя.",
        "steps": [
            "1. Открыть карточку существующей идеи.",
            "2. Проверить отображение заголовка, описания и статуса.",
            "3. Проверить доступность кнопок модерации, голосования или решения согласно роли.",
        ],
        "data": "ideaId=existing active idea, userRole=employee/admin/director",
        "expected": "Страница показывает полные сведения об идее, а набор кнопок соответствует правам текущего пользователя.",
        "actual": "Проверка не выполнялась.",
        "precondition": "В системе есть хотя бы одна идея в активном процессе.",
        "postcondition": "Пользователь остается на странице деталей идеи без потери контекста.",
        "comments": "Frontend. Для директора не должны отображаться элементы голосования, для сотрудника - элементы модерации.",
    },
    {
        "area": "Frontend",
        "priority": "Высокий",
        "title": "Создание нового сотрудника директором",
        "summary": "Проверить, что только директор видит форму добавления пользователя и может создать администратора или сотрудника.",
        "steps": [
            "1. Войти под директором.",
            "2. Открыть раздел сотрудников.",
            "3. Заполнить данные нового пользователя и выбрать роль.",
            "4. Нажать кнопку создания.",
        ],
        "data": "fullName=София Колбасенко, phone=+79990001123, password=secret123, role=admin, position=Office Administrator",
        "expected": "Новый сотрудник добавляется в список, а форма очищается или закрывается после успешного сохранения.",
        "actual": "Проверка не выполнялась.",
        "precondition": "Пользователь авторизован как директор.",
        "postcondition": "В компании появляется новая учетная запись с заданной ролью.",
        "comments": "Frontend. Проверить, что у сотрудника нет доступа к этой форме даже через прямой URL.",
    },
    {
        "area": "Frontend",
        "priority": "Средний",
        "title": "Голосование за идею и обновление результата на экране",
        "summary": "Проверить, что frontend отправляет голос и отображает обновленное состояние идеи после голосования.",
        "steps": [
            "1. Открыть идею в статусе voting.",
            "2. Нажать кнопку голоса 'за' или 'против'.",
            "3. Дождаться обновления статуса и процента поддержки.",
        ],
        "data": "voteValue=for",
        "expected": "Голос сохраняется, интерфейс показывает актуальный процент поддержки и, при необходимости, переводит идею в director_review или rejected_by_vote.",
        "actual": "Проверка не выполнялась.",
        "precondition": "Пользователь имеет право голосовать, идея находится в статусе voting.",
        "postcondition": "Голос пользователя зафиксирован, интерфейс обновлен.",
        "comments": "Frontend. Проверить блокировку повторного голосования тем же пользователем.",
    },
    {
        "area": "Backend",
        "priority": "Высокий",
        "title": "Регистрация компании и создание аккаунта директора через API",
        "summary": "Проверить, что backend создает компанию, директора и начальное состояние данных при регистрации.",
        "steps": [
            "1. Отправить POST /api/auth/register-company с валидным payload.",
            "2. Получить ответ сервера.",
            "3. Проверить сохранение новой компании в хранилище.",
        ],
        "data": "companyName=Acme, companyDescription=Internal idea platform, directorName=Иван Иванов, directorPosition=Директор, phone=+79990001122, password=secret123",
        "expected": "Сервер возвращает успешный ответ, создает компанию и директорскую учетную запись, а также подготавливает стартовые данные для дальнейшей работы.",
        "actual": "Проверка не выполнялась.",
        "precondition": "Файл хранилища доступен для записи, тестовая база пуста или содержит независимые данные.",
        "postcondition": "В системе присутствует новая компания с директором.",
        "comments": "Backend. Проверить, что повторная регистрация с тем же номером телефона отклоняется.",
    },
    {
        "area": "Backend",
        "priority": "Высокий",
        "title": "Аутентификация, получение профиля и выход из системы",
        "summary": "Проверить lifecycle токена: логин, запрос текущего пользователя и logout через backend.",
        "steps": [
            "1. Отправить POST /api/auth/login с корректными данными.",
            "2. Использовать полученный token в GET /api/auth/me.",
            "3. Отправить POST /api/auth/logout и повторить запрос me.",
        ],
        "data": "phone=+70000000003, password=employee123",
        "expected": "Сервер выдает токен, возвращает данные пользователя по me и после logout отклоняет запросы со старым токеном.",
        "actual": "Проверка не выполнялась.",
        "precondition": "В системе существует активная учетная запись сотрудника.",
        "postcondition": "Сессия завершена, токен больше не действителен.",
        "comments": "Backend. Проверить корректную ошибку при попытке входа с неверным паролем.",
    },
    {
        "area": "Backend",
        "priority": "Высокий",
        "title": "Проверка ограничений ролей при добавлении сотрудников",
        "summary": "Проверить, что только директор может создавать учетные записи сотрудников через backend API.",
        "steps": [
            "1. Отправить POST /api/employees под директором.",
            "2. Повторить запрос под администратором.",
            "3. Повторить запрос под сотрудником.",
        ],
        "data": "role=admin, phone=+79990001123, password=secret123, position=Office Administrator",
        "expected": "Запрос от директора завершается успешно, а запросы от администратора и сотрудника получают отказ в доступе.",
        "actual": "Проверка не выполнялась.",
        "precondition": "Директор авторизован, в системе есть другие роли для проверки ограничений.",
        "postcondition": "Новая учетная запись создается только в разрешенном сценарии.",
        "comments": "Backend. Проверить HTTP-коды 403 или эквивалентные ошибки доступа для неавторизованных ролей.",
    },
    {
        "area": "Backend",
        "priority": "Высокий",
        "title": "Ограничение на создание не более трех идей в месяц",
        "summary": "Проверить, что backend блокирует четвертую идею одного пользователя в пределах одного UTC-месяца.",
        "steps": [
            "1. Создать три идеи от одного пользователя.",
            "2. Отправить четвертый запрос POST /api/ideas в том же месяце.",
            "3. Проверить ответ сервера.",
        ],
        "data": "title=Idea #1..#4, description=Any valid description",
        "expected": "Первые три идеи принимаются, четвертая получает ошибку конфликта с сообщением о превышении лимита.",
        "actual": "Проверка не выполнялась.",
        "precondition": "У пользователя еще не достигнут месячный лимит идей.",
        "postcondition": "В хранилище остаются только разрешенные идеи, лимит не превышен.",
        "comments": "Backend. Проверить сброс лимита после смены UTC-месяца.",
    },
    {
        "area": "Backend",
        "priority": "Высокий",
        "title": "Модерация идеи и переход в статус voting или rejected_by_admin",
        "summary": "Проверить, что backend корректно обрабатывает решение модератора по pending-идее.",
        "steps": [
            "1. Отправить POST /api/ideas/{ideaId}/moderate для pending-идеи.",
            "2. Проверить вариант approved=true.",
            "3. Проверить вариант approved=false.",
        ],
        "data": "approved=true/false, comment=The request is clear and can go to voting.",
        "expected": "При одобрении идея переходит в voting, при отклонении - в rejected_by_admin с сохранением комментария.",
        "actual": "Проверка не выполнялась.",
        "precondition": "Идея находится в статусе pending_moderation, запрос выполняет админ или директор.",
        "postcondition": "Статус идеи изменен согласно решению модерации.",
        "comments": "Backend. Нужно убедиться, что повторная модерация уже обработанной идеи запрещена.",
    },
    {
        "area": "Backend",
        "priority": "Высокий",
        "title": "Подсчет голосов и перевод идеи в director_review или rejected_by_vote",
        "summary": "Проверить, что backend считает процент поддержки от числа eligible пользователей и корректно меняет статус идеи.",
        "steps": [
            "1. Перевести идею в статус voting.",
            "2. Отправить несколько голосов через POST /api/ideas/{ideaId}/vote.",
            "3. Проверить переход в director_review при support > 50%.",
            "4. Проверить переход в rejected_by_vote, если все голоса отданы и поддержка не превышает 50%.",
        ],
        "data": "value=for, eligibleUsers=non-director employees and admins",
        "expected": "Сервер корректно считает процент поддержки и переводит идею в нужный следующий статус без ручного вмешательства.",
        "actual": "Проверка не выполнялась.",
        "precondition": "Идея уже прошла модерацию и находится в голосовании.",
        "postcondition": "Идея либо ожидает решения директора, либо отправлена в архив после отклонения по голосам.",
        "comments": "Backend. Нужно отдельно проверить запрет голосования директором.",
    },
    {
        "area": "Backend",
        "priority": "Высокий",
        "title": "Решение директора по идее и архивирование результата",
        "summary": "Проверить, что backend принимает финальное решение директора и переводит идею в архивное состояние.",
        "steps": [
            "1. Отправить POST /api/ideas/{ideaId}/decision с approved=true.",
            "2. Повторить запрос с approved=false.",
            "3. Проверить, что идея перемещается в архив и больше не участвует в активных списках.",
        ],
        "data": "approved=true/false, comment=Rejected this quarter because of budget limits.",
        "expected": "После решения директора идея получает финальный архивный статус, а комментарий сохраняется в истории.",
        "actual": "Проверка не выполнялась.",
        "precondition": "Идея находится в статусе director_review, запрос выполняет директор.",
        "postcondition": "Идея завершена и отображается только в архиве.",
        "comments": "Backend. Директор не должен иметь возможность голосовать или создавать новые идеи.",
    },
]


def set_cell_text(cell, text):
    cell.text = text
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in paragraph.runs:
            run.font.name = "Arial"
            run.font.size = Pt(10)


def fill_metadata_table(table):
    today = datetime.now().strftime("%d.%m.%Y")
    values = [
        "Cooperate",
        "1.0",
        "Codex",
        today,
    ]
    for row, value in zip(table.rows, values):
        set_cell_text(row.cells[1], value)


def fill_case_table(table, case_index, scenario):
    rows = table.rows
    fields = [
        str(case_index),
        scenario["priority"],
        f'[{scenario["area"]}] {scenario["title"]}',
        scenario["summary"],
        "\n".join(scenario["steps"]),
        scenario["data"],
        scenario["expected"],
        scenario["actual"],
        "Не выполнен",
        scenario["precondition"],
        scenario["postcondition"],
        scenario["comments"],
    ]
    for row, value in zip(rows, fields):
        set_cell_text(row.cells[1], value)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def style_banner_table(table, text):
    table.autofit = True
    cell = table.rows[0].cells[0]
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(255, 255, 255)
    set_cell_shading(cell, "2F75B5")


def insert_banner_before(doc, source_table, text):
    banner = doc.add_table(rows=1, cols=1)
    style_banner_table(banner, text)
    source_table._tbl.addprevious(banner._tbl)
    return banner


def style_heading(paragraph, title):
    paragraph.clear()
    run = paragraph.add_run(title)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(46, 116, 181)
    paragraph.style = "Heading 2"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)


def find_paragraph_by_text(doc, text):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    raise ValueError(f'Paragraph not found: {text}')


def add_page_break(doc):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    run.add_break(WD_BREAK.PAGE)


def clone_table(doc, source_table):
    new_tbl = deepcopy(source_table._tbl)
    doc._body._element.append(new_tbl)
    return doc.tables[-1]


def main():
    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(f"Template not found: {TEMPLATE_PATH}")

    doc = Document(str(TEMPLATE_PATH))

    fill_metadata_table(doc.tables[1])

    base_heading = find_paragraph_by_text(doc, "Тестовый пример #1:")
    base_table = doc.tables[3]

    intro_break = base_heading.insert_paragraph_before()
    intro_break.add_run().add_break(WD_BREAK.PAGE)
    style_heading(base_heading, "Тестовый пример #1:")
    insert_banner_before(doc, base_table, "Блок тестового сценария 1 - Frontend")
    fill_case_table(base_table, 1, SCENARIOS[0])
    add_page_break(doc)

    for index, scenario in enumerate(SCENARIOS[1:], start=2):
        heading = doc.add_paragraph()
        heading.style = base_heading.style
        style_heading(heading, f"Тестовый пример #{index}:")
        table = clone_table(doc, base_table)
        insert_banner_before(doc, table, f'Блок тестового сценария {index} - {scenario["area"]}')
        fill_case_table(table, index, scenario)
        if index != len(SCENARIOS):
            add_page_break(doc)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT_PATH))
    print(f"Saved {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
